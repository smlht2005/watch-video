#!/usr/bin/env python3
"""Build a readable report (<stem>_report.md + <stem>_report.docx) from an archived watch folder.

Reads  : meta.json, summary.md, transcript.md, frames.json, frames/*.jpg,
         transcript.zh.md (optional: Traditional Chinese translation of a foreign-language transcript,
         same "[mm:ss] text" lines; shown under the original text of each section),
         supplement.md (optional, hand-written, never overwritten by watch_archive: information the video
         does not show, e.g. the full text of a screen capture that was cut off, post captions, links)
Writes : <stem>_report.md   (frames interleaved with the transcript)
         <stem>_report.docx (same content as a Word document)

Usage: python3 scripts/build_report.py watch-output/<folder>
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

CJK_FONT = "Microsoft JhengHei"
SEG_RE = re.compile(r"^\[(\d+:\d+(?::\d+)?)\]\s*(.*)$")


def file_stem(dest: Path) -> str:
    """<date>_<platform>_<title (<=20 chars)>__<video-id>, derived from the folder name."""
    head, _, vid = dest.name.rpartition("__")
    date, platform, title = (head.split("_", 2) + ["", ""])[:3]
    return f"{date}_{platform}_{title[:20].rstrip('-')}__{vid}"


def to_sec(ts: str) -> int:
    n = [int(x) for x in ts.split(":")]
    while len(n) < 3:
        n.insert(0, 0)
    return n[0] * 3600 + n[1] * 60 + n[2]


def read_segs(path: Path):
    segs = []
    for line in path.read_text(encoding="utf-8").splitlines():
        m = SEG_RE.match(line.strip())
        if m:
            segs.append((m.group(1), m.group(2)))
    return segs


def load(dest: Path):
    meta = json.loads((dest / "meta.json").read_text(encoding="utf-8"))
    summary = (dest / "summary.md").read_text(encoding="utf-8")
    segs = read_segs(dest / "transcript.md")
    zpath = dest / "transcript.zh.md"
    meta["_zh"] = read_segs(zpath) if zpath.exists() else []
    fpath = dest / "frames.json"
    frames = json.loads(fpath.read_text(encoding="utf-8")) if fpath.exists() else []
    return meta, summary, segs, frames


def join_text(parts):
    """Join transcript fragments: spaces for Latin-script text, Chinese commas otherwise."""
    parts = [p.strip() for p in parts if p and p.strip()]
    if not parts:
        return ""
    body = "".join(parts)
    if sum(c.isascii() for c in body) > 0.6 * len(body):
        text = " ".join(parts)
        return text if text[-1] in ".?!" else text + "."
    text = "".join(p if p[-1] in "，。？！；：、" else p + "，" for p in parts)
    return text.rstrip("，") + ("" if text.rstrip("，")[-1] in "。？！" else "。")


def text_lines(meta, segs, lo, hi):
    """Original text for [lo, hi) seconds, plus its translation from transcript.zh.md if present."""
    L = []
    text = join_text(t for ts, t in segs if lo <= to_sec(ts) < hi)
    if text:
        L += [text, ""]
    zh = "".join(t.strip() for ts, t in meta.get("_zh", []) if lo <= to_sec(ts) < hi)
    if zh:
        L += [f"**中譯：** {zh}", ""]
    return L


def sections(segs, frames):
    """Group transcript segments under the frame whose time window contains them."""
    if not frames:
        return [(None, segs)]
    frames = sorted(frames, key=lambda f: to_sec(f["t"]))
    out = [(f, []) for f in frames]
    for ts, text in segs:
        s = to_sec(ts)
        idx = 0
        for i, f in enumerate(frames):  # first frame among those sharing the latest start time
            if to_sec(f["t"]) <= s and to_sec(f["t"]) > to_sec(frames[idx]["t"]) or i == 0:
                idx = i
        out[idx][1].append((ts, text))
    return out


def strip_h1(md: str) -> str:
    return re.sub(r"\A# .*\n+", "", md)


def slide_blocks(meta, segs):
    """One block per slide: one frame, optional note, then the transcript spoken while it is on screen."""
    slides = sorted(meta["slides"], key=lambda x: to_sec(x["start"]))
    end = meta.get("duration", "").split(" ")[0]
    L = []
    for n, sl in enumerate(slides):
        nxt = slides[n + 1]["start"] if n + 1 < len(slides) else end
        lo, hi = to_sec(sl["start"]), (to_sec(nxt) if nxt else 10 ** 9)
        title = f"：{sl['title']}" if sl.get("title") else ""
        L += [f"### 投影片 {n + 1}{title}（{sl['start']}–{nxt}）", "", f"![{sl['start']}](frames/{sl['frame']})", ""]
        if sl.get("note"):
            L += [f"**重點：** {sl['note']}", ""]
        if sl.get("keywords"):
            L += ["**標籤：** " + " ".join(f"#{k}" for k in sl["keywords"]), ""]
        L += text_lines(meta, segs, lo, hi)
    return L


def frame_blocks(meta, segs, frames):
    secs = [(f, it) for f, it in sections(segs, frames) if f is None or it]  # skip frames with no speech
    end = meta.get("duration", "").split(" ")[0]
    L = []
    for n, (frame, items) in enumerate(secs):
        nxt = secs[n + 1][0]["t"] if frame and n + 1 < len(secs) else end
        if frame:
            L += [f"### 畫面 {frame['t']}–{nxt}", "", f"![{frame['t']}](frames/{frame['file']})", ""]
            tags = meta.get("section_keywords", {}).get(frame["t"], [])
            if tags:
                L += ["**標籤：** " + " ".join(f"#{k}" for k in tags), ""]
        lo = to_sec(frame["t"]) if frame and n else 0
        hi = to_sec(nxt) if frame and nxt else 10 ** 9
        L += text_lines(meta, segs, lo, hi)
    return L


FENCE_OPEN = re.compile(r"^\s*(`{3,}|~{3,})")


def fence_close(marker: str, line: str) -> bool:
    """A closing fence is a bare run of the same character, at least as long as the opening one (so ```python never closes)."""
    m = re.fullmatch(r"\s*(`{3,}|~{3,})\s*", line)
    return bool(m) and m.group(1)[0] == marker[0] and len(m.group(1)) >= len(marker)


def close_open_fence(md: str) -> str:
    """Append a closing fence if the markdown ends inside one, so a stray ``` cannot swallow the rest of the report."""
    fence = None
    for line in md.splitlines():
        if fence:
            if fence_close(fence, line):
                fence = None
        elif (m := FENCE_OPEN.match(line)):
            fence = m.group(1)
    return md + "\n" + fence if fence else md


FRAME_REF = re.compile(r"!\[([^\]]*)\]\(frame@(\d+:\d+(?::\d+)?)\)")


def resolve_frame_ref(line: str, frames: list) -> str:
    """Replace ![alt](frame@MM:SS) with the frames/ file closest in time (frames/ is renumbered on every run, so
    hand-written files refer to a time, not a file name). With no frames, leave a text note instead of a broken image."""
    def repl(m):
        cands = [(abs(to_sec(f["t"]) - to_sec(m.group(2))), f["file"]) for f in frames if f.get("file") and f.get("t")]
        if not cands:
            return f"（缺截圖：{m.group(2)}）"
        return f"![{m.group(1)}](frames/{min(cands)[1]})"
    return FRAME_REF.sub(repl, line)


def include_lines(path: Path, heading: str, frames=None):
    """Body of a hand-written side file (supplement.md, sop.md) for the report, under '## <heading>'.

    The file's leading '# …' title is dropped; '#'/'##' headings -> '###' heading; '###' and deeper -> bold line
    (the docx converter handles at most three heading levels). Fenced blocks (``` or ~~~) are kept verbatim with
    their original markers, so a ~~~ block may contain ```; an unclosed fence is closed at the end, and md_to_docx
    renders fences as indented monospace. Outside fences, ![alt](frame@MM:SS) is resolved against `frames` (when
    given). Returns [] when the file is missing or has no content besides its title.
    """
    if not path.exists():
        return []
    out, fence, has_body, started = [], None, False, False
    for line in path.read_text(encoding="utf-8").splitlines():
        m_fence = FENCE_OPEN.match(line)
        if fence:
            if fence_close(fence, line):
                out.append(fence)
                fence = None
            else:
                out.append(line)
        elif m_fence:
            fence, has_body, started = m_fence.group(1), True, True
            out.append(fence)
        elif (m := re.match(r"^(#{1,6})\s+(.*)$", line)):
            level, text = len(m.group(1)), m.group(2).strip()
            if level == 1 and not started:  # the file's own title; the section heading replaces it
                started = True
                continue
            started = has_body = True
            out.append(f"**{text}**" if level >= 3 else f"### {text}")
        else:
            if line.strip():
                started = has_body = True
            out.append(resolve_frame_ref(line, frames) if frames is not None else line)
    if fence:
        out.append(fence)
    if not has_body:
        return []
    return [f"## {heading}", ""] + out + [""]


def supplement_lines(dest: Path):
    return include_lines(dest / "supplement.md", "補充資料（影片未呈現的資訊）")


def sop_lines(dest: Path, frames: list):
    return include_lines(dest / "sop.md", "操作 SOP（依影片整理，照著做）", frames)


def build_md(dest: Path, meta, summary, segs, frames) -> str:
    L = [f"# {meta['title']}", ""]
    L += [f"- 來源：{meta['url']}",
          f"- 平台 / 上傳者：{meta['platform']} / {meta.get('uploader', '')}",
          f"- 長度：{meta.get('duration', '')}",
          f"- 逐字稿來源：{meta.get('transcript_source', '')}",
          f"- 更新日期：{meta.get('last_updated', '')}"]
    if meta.get("keywords"):
        L.append("- **關鍵字：** " + " ".join(f"#{k}" for k in meta["keywords"]))
    L.append("")
    L += ["## 摘要與名詞說明", "", close_open_fence(strip_h1(summary).strip()), ""]
    if meta.get("sop_mode") != "off":
        L += sop_lines(dest, frames)
    sup = supplement_lines(dest)
    L += sup
    L += ["## 圖文對照（畫面 + 逐字稿" + ("＋中譯）" if meta.get("_zh") else "）"), ""]
    if meta.get("slides"):
        L += slide_blocks(meta, segs)
    else:
        L += frame_blocks(meta, segs, frames)
    if meta.get("known_asr_errors"):
        L += ["## 已知轉錄錯誤", "", "| 時間 | 聽成 | 應為 |", "|---|---|---|"]
        for e in meta["known_asr_errors"]:
            L.append(f"| {e['time']} | {e['heard']} | {e['expected']} |")
        L.append("")
    if sup:
        L += ["## 補充資料連結", "",
              "- 補充資料檔：[supplement.md](supplement.md)（與本報告同資料夾；內容已收錄於上方〈補充資料〉）", ""]
    return "\n".join(L)


def sop_md(dest: Path, meta, frames):
    """Standalone SOP document (no transcript): title, source, then the SOP chapters one heading level up."""
    if meta.get("sop_mode") == "off":
        return None
    lines = sop_lines(dest, frames)
    if not lines:
        return None
    body = ["## " + l[4:] if l.startswith("### ") else l for l in lines[2:]]
    head = [f"# {meta['title']}：操作 SOP", "",
            f"- 來源：{meta['url']}（步驟後標註影片時間）",
            f"- 更新日期：{meta.get('last_updated', '')}",
            "- 說明：本文由影片整理；影片中的價格、硬體需求與效果為講者主張，未查證。", ""]
    return "\n".join(head + body)


def sop_name(meta) -> str:
    """Readable file stem for the SOP outputs: '<video title>_操作SOP' (characters Windows forbids removed)."""
    safe = re.sub(r'[\\/:*?"<>|]', "", meta["title"]).strip()[:60].strip() or "video"
    return f"{safe}_操作SOP"


def build_kb(dest: Path, meta, frames):
    """Portable knowledge-base bundle for Obsidian/Notion: kb/<title>.md (YAML frontmatter) + kb/attachments/.

    Images get stable names from their video time (step-MMSS.jpg) and relative links, so the folder can be copied
    anywhere and survives the renumbering of frames/ on a rerun. Returns the .md path, or None when there is no SOP.
    """
    import shutil
    kb = dest / "kb"
    shutil.rmtree(kb, ignore_errors=True)
    body = sop_md(dest, meta, frames)
    if not body:
        return None
    t_of = {f["file"]: f["t"] for f in frames if f.get("file") and f.get("t")}
    (kb / "attachments").mkdir(parents=True)

    def repl(m):
        src = dest / "frames" / m.group(2)
        if not src.exists():
            return m.group(0)
        name = "step-" + t_of.get(m.group(2), m.group(2).rsplit(".", 1)[0]).replace(":", "") + ".jpg"
        shutil.copyfile(src, kb / "attachments" / name)
        return f"![{m.group(1)}](attachments/{name})"

    body = re.sub(r"!\[([^\]]*)\]\(frames/([^)]+)\)", repl, body)
    tags = ["sop", "tutorial"] + [re.sub(r"\s+", "-", k.strip()) for k in meta.get("keywords", []) if k.strip()]
    title = meta["title"].replace('"', "'")
    front = ["---", f'title: "{title}：操作 SOP"', "type: sop", f"source: {meta['url']}",
             f"platform: {meta.get('platform', '')}", f"video_id: {meta.get('video_id', '')}",
             f"date: {meta.get('last_updated', '')}", "tags:"] + [f"  - {t}" for t in tags] + ["---", ""]
    out = kb / f"{sop_name(meta)}.md"
    out.write_text("\n".join(front) + body + "\n", encoding="utf-8")
    return out


# ---------- Markdown -> docx (small subset: headings, tables, bullets, images, **bold**, `code`) ----------

def set_font(run, size=None, bold=None, mono=False):
    name = "Consolas" if mono else CJK_FONT
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


URL_RE = re.compile(r"https?://[A-Za-z0-9._~:/?#@!$&'*+,;=%-]+")


def add_hyperlink(par, url, size, label=None):
    """Append a real clickable hyperlink run to a python-docx paragraph."""
    rid = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, attrs in (("w:rStyle", {"w:val": "Hyperlink"}), ("w:color", {"w:val": "0563C1"}), ("w:u", {"w:val": "single"}),
                       ("w:sz", {"w:val": str(int(size * 2))})):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        rpr.append(el)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = label or url
    run.append(text)
    link.append(run)
    par._p.append(link)


def add_inline(par, text, size=10.5):
    for tok in re.split(r"(\[[^\]]+\]\((?:[^()\s]|\([^()\s]*\))+\)|\*\*[^*]+\*\*|`[^`]+`|https?://[A-Za-z0-9._~:/?#@!$&'*+,;=%-]+)", text):
        if not tok:
            continue
        if (lm := re.fullmatch(r"\[([^\]]+)\]\(((?:[^()\s]|\([^()\s]*\))+)\)", tok)):
            label, target = lm.groups()
            if URL_RE.fullmatch(target):
                add_hyperlink(par, target, size, label)
            else:
                set_font(par.add_run(label), size, mono=True)
        elif URL_RE.fullmatch(tok):
            add_hyperlink(par, tok, size)
        elif tok.startswith("**"):
            set_font(par.add_run(tok[2:-2]), size, bold=True)
        elif tok.startswith("`"):
            set_font(par.add_run(tok[1:-1]), size - 1, mono=True)
        else:
            set_font(par.add_run(tok), size)


def crop_dark_borders(im, threshold=16, min_trim=0.05):
    """Trim near-black letterbox / pillarbox bars (e.g. a notebook recorded in the middle of a wide frame).

    Only bars wider than min_trim of the image are cut, so normal dark-edged frames stay as they are. The saved frame
    files are not changed; this only affects the picture inserted into the docx.
    """
    gray = im.convert("L")
    w, h = gray.size
    px = gray.load()
    step_y, step_x = max(1, h // 60), max(1, w // 60)

    def col_dark(x):
        return max(px[x, y] for y in range(0, h, step_y)) <= threshold

    def row_dark(y):
        return max(px[x, y] for x in range(0, w, step_x)) <= threshold

    left = next((x for x in range(w) if not col_dark(x)), 0)
    right = next((x for x in range(w - 1, -1, -1) if not col_dark(x)), w - 1) + 1
    top = next((y for y in range(h) if not row_dark(y)), 0)
    bottom = next((y for y in range(h - 1, -1, -1) if not row_dark(y)), h - 1) + 1
    if right <= left or bottom <= top:
        return im  # an all-dark frame: leave it alone
    left = left if left >= w * min_trim else 0
    right = right if w - right >= w * min_trim else w
    top = top if top >= h * min_trim else 0
    bottom = bottom if h - bottom >= h * min_trim else h
    return im.crop((left, top, right, bottom)) if (left, top, right, bottom) != (0, 0, w, h) else im


def md_to_docx(dest: Path, md: str, out: Path, kw=None, max_img_h=9):
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    for name, size in (("Normal", 10.5), ("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12)):
        st = doc.styles[name]
        st.font.name = CJK_FONT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        st.font.size = Pt(size)

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if (fm := FENCE_OPEN.match(line)):  # fenced block (``` or ~~~): verbatim, indented, monospace
            marker = fm.group(1)
            i += 1
            while i < len(lines) and not fence_close(marker, lines[i]):
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.8)
                par.paragraph_format.space_after = Pt(0)
                set_font(par.add_run(lines[i]), 10, mono=True)
                i += 1
            i += 1
            doc.add_paragraph()
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol)
            t.style = "Table Grid"
            for r, row in enumerate(rows):
                for c in range(ncol):
                    cell = t.cell(r, c)
                    cell.text = ""
                    par = cell.paragraphs[0]
                    add_inline(par, row[c] if c < len(row) else "", 9.5)
                    if r == 0:
                        for run in par.runs:
                            run.bold = True
            doc.add_paragraph()
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
        elif (im := re.match(r"^!\[[^\]]*\]\(([^)]+)\)$", line.strip())):
            img = dest / im.group(1)
            if img.exists():
                from PIL import Image
                import io
                pil = crop_dark_borders(Image.open(img).convert("RGB"))
                w, h = pil.size
                # portrait (vertical) videos: cap height at 7 cm; landscape/screen recordings: up to max_img_h cm tall, 16 cm wide
                width_cm = min(12, 7 * w / h) if h > w else min(16, max_img_h * w / h)
                buf = io.BytesIO()  # re-encode: some ffmpeg JPEGs lack a JFIF header that python-docx requires
                pil.save(buf, "JPEG", quality=90)
                buf.seek(0)
                doc.add_picture(buf, width=Cm(width_cm))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.lstrip().startswith("- "):  # also indented sub-items under a numbered step
            add_inline(doc.add_paragraph(style="List Bullet"), line.lstrip()[2:])
        elif line.strip():
            add_inline(doc.add_paragraph(), line.strip())
        i += 1
    if kw:
        doc.core_properties.keywords = ", ".join(kw)
    doc.save(str(out))


def build(dest: Path):
    meta, summary, segs, frames = load(dest)
    md = build_md(dest, meta, summary, segs, frames)
    stem = file_stem(dest)
    for old in ("report.md", "report.docx", *dest.glob("*_report.*"), *dest.glob("*_sop.*"), *dest.glob("*_操作SOP.*")):  # drop stale / legacy names
        (old if isinstance(old, Path) else dest / old).unlink(missing_ok=True)  # glob() already includes dest
    out_md, out_docx = dest / f"{stem}_report.md", dest / f"{stem}_report.docx"
    out_md.write_text(md, encoding="utf-8")
    md_to_docx(dest, md, out_docx, meta.get("keywords"))
    outs = [out_md, out_docx]
    standalone = sop_md(dest, meta, frames)
    if standalone:
        s_md, s_docx = dest / f"{sop_name(meta)}.md", dest / f"{sop_name(meta)}.docx"
        s_md.write_text(standalone, encoding="utf-8")
        md_to_docx(dest, standalone, s_docx, meta.get("keywords"), max_img_h=6.5)  # many step shots: keep them compact
        outs += [s_md, s_docx]
    build_kb(dest, meta, frames)  # also clears a stale kb/ when there is no SOP
    return tuple(outs)


if __name__ == "__main__":
    for p in build(Path(sys.argv[1])):
        print("wrote", p)
