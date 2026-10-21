"""Unit tests for language detection and supplement.md handling. Run: python3 -m unittest discover -s tests -v"""
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import build_report as br  # noqa: E402
import pick_frames as pf  # noqa: E402
import review_links as rl  # noqa: E402
import watch_archive as wa  # noqa: E402


class DetectLanguage(unittest.TestCase):
    def test_english(self):
        self.assertEqual(wa.detect_language("[00:01] Welcome to this tutorial on Kind and how clusters work today.\n" * 3), "en")

    def test_chinese(self):
        self.assertEqual(wa.detect_language("[00:01] 如果你的額度常常不到一天就用完了，教你一個隱藏用法。\n" * 3), "zh")

    def test_chinese_with_many_english_terms(self):
        text = "[00:01] 用 kubectl config get-contexts 查看 cluster 與 context，再用 kind create cluster 建立。\n" * 3
        self.assertEqual(wa.detect_language(text), "zh")

    def test_japanese_not_zh(self):
        self.assertEqual(wa.detect_language("[00:01] これはテストです。今日は東京で会議があります。\n" * 5), "ja")

    def test_korean(self):
        self.assertEqual(wa.detect_language("[00:01] 안녕하세요 오늘은 서울에서 회의가 있습니다 감사합니다\n" * 5), "ko")

    def test_too_short_is_unknown(self):
        self.assertEqual(wa.detect_language("[00:00] hi"), "unknown")
        self.assertEqual(wa.detect_language(""), "unknown")

    def test_ignores_timestamps_and_headers(self):
        text = "# 標題標題標題標題標題標題標題標題\n_Source: whisper (groq)._\n```\n[00:01] ok\n```\n"
        self.assertEqual(wa.detect_language(text), "unknown")

    def test_real_japanese_prose_is_ja(self):
        self.assertEqual(wa.detect_language("[00:01] これはテストです。今日は東京で会議があります。\n" * 5), "ja")

    def test_kanji_heavy_japanese_is_left_for_human_check(self):
        # hiragana is only ~9% of the CJK characters: not confident enough for 'ja', must not become 'zh'
        text = "[00:01] 日本国内経済成長率予測発表政府統計調査結果概要説明資料公開である\n" * 3
        self.assertEqual(wa.detect_language(text), "other")

    CHINESE_LINES = ("[00:01] 今天要跟大家分享一個很實用的方法，可以讓你的工作效率提升很多，而且完全不用花錢。\n"
                     "[00:08] 首先打開設定頁面，選擇你需要的功能，接著按照畫面上的步驟一步一步操作就可以了。\n"
                     "[00:15] 最後記得儲存，這樣下次打開的時候就會自動套用，非常方便，大家可以試試看。\n")

    def test_chinese_with_one_japanese_phrase_is_zh(self):
        text = self.CHINESE_LINES + "[00:22] 影片最後講者用日文說了ありがとうございます，謝謝大家的收看。\n"
        self.assertEqual(wa.detect_language(text), "zh")

    def test_chinese_mentioning_japanese_titles_is_zh(self):
        text = self.CHINESE_LINES + "[00:22] 我最喜歡的宮崎駿作品是となりのトトロ和もののけ姫，劇情非常感人，畫面也很美。\n"
        self.assertEqual(wa.detect_language(text), "zh")

    def test_chinese_with_katakana_brands_is_zh(self):
        self.assertEqual(wa.detect_language("[00:01] 今天我們去了東京的ユニクロ和セブンイレブン，買了很多東西回家。\n" * 3), "zh")
        self.assertEqual(wa.detect_language("[00:01] 這台車的外觀很好看，內裝比トヨタ還要精緻，但是價格比トヨタ貴很多。\n" * 3), "zh")

    def test_short_real_chinese_is_unknown_not_other(self):
        self.assertEqual(wa.detect_language("[00:01] 好的謝謝\n"), "unknown")

    def test_repeated_single_word_is_unknown(self):
        self.assertEqual(wa.detect_language("[00:01] ok\n" * 15), "unknown")

    def test_mostly_english_with_little_chinese_is_other(self):
        text = "[00:01] hello world this is some english words and more english text here 中文\n" * 5
        self.assertEqual(wa.detect_language(text), "other")  # too much Chinese for 'en', too little for 'zh'

    def test_balanced_mixed_text_is_zh_by_design(self):
        text = "[00:01] hello world this is some english words and 中文字幾個\n" * 3
        self.assertEqual(wa.detect_language(text), "zh")  # >=10 han characters and >=20% of the units


class LanguagePolicy(unittest.TestCase):
    def test_translation_flag_only_for_en_ja_ko(self):
        for lang, expected in (("en", True), ("ja", True), ("ko", True), ("zh", False), ("other", False), ("unknown", False)):
            self.assertEqual(wa.needs_zh_translation(lang), expected, lang)

    def test_opencc_skipped_for_ja_ko_other(self):
        for lang, expected in (("zh", True), ("unknown", True), ("en", True), ("ja", False), ("ko", False), ("other", False)):
            self.assertEqual(wa.convert_to_traditional(lang), expected, lang)

    def test_opencc_would_really_corrupt_japanese(self):
        if wa._S2T is None:
            self.skipTest("opencc not installed")
        self.assertNotEqual(wa.to_traditional("会議学校"), "会議学校")  # the reason for the ja/ko skip


class CropDarkBorders(unittest.TestCase):
    def test_pillarboxed_frame_is_trimmed(self):
        from PIL import Image
        im = Image.new("RGB", (1024, 576), "black")
        im.paste(Image.new("RGB", (512, 576), "white"), (256, 0))
        self.assertEqual(br.crop_dark_borders(im).size, (512, 576))

    def test_normal_and_all_dark_frames_untouched(self):
        from PIL import Image
        normal = Image.new("RGB", (640, 360), "gray")
        self.assertEqual(br.crop_dark_borders(normal).size, (640, 360))
        dark = Image.new("RGB", (640, 360), "black")
        self.assertEqual(br.crop_dark_borders(dark).size, (640, 360))

    def test_thin_dark_edge_is_kept(self):
        from PIL import Image
        im = Image.new("RGB", (1000, 500), "white")
        im.paste(Image.new("RGB", (20, 500), "black"), (0, 0))  # 2% bar: below the 5% minimum
        self.assertEqual(br.crop_dark_borders(im).size, (1000, 500))


class RemapSlideFrames(unittest.TestCase):
    def test_slides_follow_their_old_frame_time_after_renumbering(self):
        old = [{"file": "frame_0002.jpg", "t": "02:05"}, {"file": "frame_0013.jpg", "t": "08:57"}]
        new = [{"file": "frame_0002.jpg", "t": "00:02"}, {"file": "frame_0007.jpg", "t": "02:05"},
               {"file": "frame_0020.jpg", "t": "08:55"}, {"file": "frame_0021.jpg", "t": "09:30"}]
        slides = [{"start": "01:43", "frame": "frame_0002.jpg"}, {"start": "03:13", "frame": "frame_0013.jpg"}]
        out = wa.remap_slide_frames(slides, old, new)
        self.assertEqual([s["frame"] for s in out], ["frame_0007.jpg", "frame_0020.jpg"])

    def test_unknown_old_frame_falls_back_to_slide_start(self):
        out = wa.remap_slide_frames([{"start": "09:35", "frame": "gone.jpg"}], [], [{"file": "a.jpg", "t": "09:30"}, {"file": "b.jpg", "t": "20:00"}])
        self.assertEqual(out[0]["frame"], "a.jpg")

    def test_hour_long_times(self):
        out = wa.remap_slide_frames([{"start": "63:58", "frame": "x.jpg"}], [], [{"file": "a.jpg", "t": "1:04:22"}, {"file": "b.jpg", "t": "10:00"}])
        self.assertEqual(out[0]["frame"], "a.jpg")


class SupplementLines(unittest.TestCase):
    def run_sup(self, text):
        d = Path(tempfile.mkdtemp())
        if text is not None:
            (d / "supplement.md").write_text(text, encoding="utf-8")
        return br.supplement_lines(d)

    def test_missing_empty_and_title_only(self):
        self.assertEqual(self.run_sup(None), [])
        self.assertEqual(self.run_sup(""), [])
        self.assertEqual(self.run_sup("# 補充資料\n\n"), [])

    def test_headings_and_fence(self):
        out = self.run_sup("# 補充資料\n\n## S1｜x\n- a\n### 原文\n```\n# in fence\n**b**\n```\n#### deep\ntext\n")
        self.assertEqual(out[0], "## 補充資料（影片未呈現的資訊）")
        self.assertIn("### S1｜x", out)
        self.assertIn("**原文**", out)
        self.assertIn("**deep**", out)
        i = out.index("```")
        self.assertEqual(out[i + 1:i + 3], ["# in fence", "**b**"])  # verbatim, not parsed as heading/bold

    def test_tilde_fence_kept_and_unclosed_fence_closed(self):
        out = self.run_sup("## S1\n~~~\ncode\n")
        self.assertEqual(out.count("~~~"), 2)  # opened as ~~~, closed with ~~~ at the end

    def test_tilde_fence_may_contain_backtick_fence(self):
        out = self.run_sup("## S1\n~~~\n```\ninner\n```\n~~~\nafter\n")
        i = out.index("~~~")
        self.assertEqual(out[i + 1:i + 4], ["```", "inner", "```"])  # inner fence does not close the outer one
        self.assertEqual(out[i + 4], "~~~")
        self.assertIn("after", out)

    def test_first_h1_dropped_even_if_not_named_supplement(self):
        out = self.run_sup("# 別的標題\n\n## S1\nbody\n")
        self.assertNotIn("# 別的標題", "\n".join(out))
        self.assertIn("### S1", out)


class SopLines(unittest.TestCase):
    FRAMES = [{"file": "cue_0001.jpg", "t": "01:00"}, {"file": "cue_0002.jpg", "t": "05:26"},
              {"file": "cue_0003.jpg", "t": "1:02:03"}]

    def run_sop(self, text, frames=None):
        d = Path(tempfile.mkdtemp())
        if text is not None:
            (d / "sop.md").write_text(text, encoding="utf-8")
        return br.sop_lines(d, self.FRAMES if frames is None else frames)

    def test_missing_empty_and_title_only(self):
        self.assertEqual(self.run_sop(None), [])
        self.assertEqual(self.run_sop(""), [])
        self.assertEqual(self.run_sop("# 操作 SOP\n\n"), [])

    def test_heading_and_chapters(self):
        out = self.run_sop("# 操作 SOP\n\n## 一、安裝\n1. 下載\n### 注意事項\n- 費用\n")
        self.assertEqual(out[0], "## 操作 SOP（依影片整理，照著做）")
        self.assertIn("### 一、安裝", out)
        self.assertIn("1. 下載", out)
        self.assertIn("**注意事項**", out)

    def test_frame_ref_picks_nearest_frame(self):
        out = "\n".join(self.run_sop("## 一\n![設定畫面](frame@05:00)\n![x](frame@1:02:00)\n"))
        self.assertIn("![設定畫面](frames/cue_0002.jpg)", out)  # 05:00 is closer to 05:26 than 01:00
        self.assertIn("![x](frames/cue_0003.jpg)", out)

    def test_frame_ref_inside_fence_untouched(self):
        out = self.run_sop("## 一\n```\n![a](frame@01:00)\n```\n")
        self.assertIn("![a](frame@01:00)", out)

    def test_frame_ref_without_any_frames_leaves_note_not_broken_image(self):
        out = "\n".join(self.run_sop("## 一\n![a](frame@01:00)\n步驟\n", frames=[]))
        self.assertNotIn("frame@", out)
        self.assertNotIn("![", out)
        self.assertIn("（缺截圖：01:00）", out)
        self.assertIn("步驟", out)

    def test_supplement_behaviour_unchanged(self):
        d = Path(tempfile.mkdtemp())
        (d / "supplement.md").write_text("# 補充資料\n\n## S1｜x\n- a\n", encoding="utf-8")
        out = br.supplement_lines(d)
        self.assertEqual(out[0], "## 補充資料（影片未呈現的資訊）")
        self.assertIn("### S1｜x", out)


class BuildWithSop(unittest.TestCase):
    SOP = "# 操作 SOP\n\n## 一、安裝\n1. 下載（00:30）\n![畫面](frame@00:31)\n"

    def make(self, sop=None, sop_mode=None):
        import json
        from PIL import Image
        root = Path(tempfile.mkdtemp())
        d = root / "2026-10-13_YouTube_測試影片__abc123"
        (d / "frames").mkdir(parents=True)
        Image.new("RGB", (320, 180), (200, 200, 200)).save(d / "frames" / "cue_0001.jpg")
        meta = {"url": "https://example.com/v", "platform": "YouTube", "title": "測試影片", "duration": "01:00 (60s)",
                "keywords": ["k"], "slides": []}
        if sop_mode:
            meta["sop_mode"] = sop_mode
        (d / "meta.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
        (d / "summary.md").write_text("# 摘要\n\n內容\n", encoding="utf-8")
        (d / "transcript.md").write_text("# t\n\n```\n[00:01] 你好\n```\n", encoding="utf-8")
        (d / "frames.json").write_text(json.dumps([{"file": "cue_0001.jpg", "t": "00:30"}]), encoding="utf-8")
        if sop is not None:
            (d / "sop.md").write_text(sop, encoding="utf-8")
        return d

    def names(self, d, suffix):
        return sorted(d / p.name for p in d.iterdir() if p.is_file() and p.name.endswith(suffix))

    def test_sop_in_report_before_supplement_and_standalone_files(self):
        d = self.make(self.SOP)
        (d / "supplement.md").write_text("## S1\n- a\n", encoding="utf-8")
        outs = br.build(d)
        self.assertEqual(len(outs), 4)
        md = outs[0].read_text(encoding="utf-8")
        self.assertLess(md.index("## 摘要與名詞說明"), md.index("## 操作 SOP"))
        self.assertLess(md.index("## 操作 SOP"), md.index("## 補充資料"))
        self.assertIn("![畫面](frames/cue_0001.jpg)", md)
        sop = self.names(d, "_操作SOP.md")[0].read_text(encoding="utf-8")
        self.assertTrue(sop.startswith("# 測試影片：操作 SOP"))
        self.assertNotIn("## 摘要", sop)
        self.assertNotIn("你好", sop)  # no transcript in the standalone SOP
        self.assertIn("## 一、安裝", sop)
        self.assertNotIn("### 一、安裝", sop)  # chapters promoted one level
        from docx import Document
        doc = Document(str(self.names(d, "_操作SOP.docx")[0]))
        self.assertEqual(len(doc.inline_shapes), 1)

    def test_no_sop_file_means_no_section_and_no_sop_outputs(self):
        d = self.make(None)
        outs = br.build(d)
        self.assertEqual(len(outs), 2)
        self.assertNotIn("操作 SOP", outs[0].read_text(encoding="utf-8"))
        self.assertEqual(self.names(d, "_操作SOP.md") + self.names(d, "_操作SOP.docx"), [])

    def test_title_only_sop_means_nothing(self):
        d = self.make("# 操作 SOP\n\n")
        self.assertEqual(len(br.build(d)), 2)

    def test_sop_mode_off_skips_even_with_file(self):
        d = self.make(self.SOP, sop_mode="off")
        outs = br.build(d)
        self.assertEqual(len(outs), 2)
        self.assertNotIn("操作 SOP", outs[0].read_text(encoding="utf-8"))

    def test_stale_sop_outputs_removed_when_sop_deleted(self):
        d = self.make(self.SOP)
        br.build(d)
        (d / "sop.md").unlink()
        br.build(d)
        self.assertEqual(self.names(d, "_操作SOP.md") + self.names(d, "_操作SOP.docx"), [])


class StaleCleanupWithRelativeDest(BuildWithSop):
    def test_legacy_outputs_removed_when_dest_is_a_relative_path(self):
        d = self.make(self.SOP)
        (d / "x_sop.md").write_text("old", encoding="utf-8")
        (d / "y_report.md").write_text("old", encoding="utf-8")
        cwd = os.getcwd()
        os.chdir(d.parent)
        try:
            br.build(Path(d.name))  # as the CLI does: watch-output/<folder>
        finally:
            os.chdir(cwd)
        self.assertFalse((d / "x_sop.md").exists())
        self.assertFalse((d / "y_report.md").exists())


class KbBundle(BuildWithSop):
    def kb_files(self, d):
        kb = d / "kb"
        return sorted(p.name for p in kb.glob("*.md")) if kb.exists() else []

    def test_bundle_has_frontmatter_portable_images_and_no_build_refs(self):
        d = self.make(self.SOP)
        br.build(d)
        self.assertEqual(self.kb_files(d), ["測試影片_操作SOP.md"])
        md = (d / "kb" / "測試影片_操作SOP.md").read_text(encoding="utf-8")
        self.assertTrue(md.startswith("---\n"))
        head = md.split("---\n")[1]
        for key in ('title: "測試影片：操作 SOP"', "type: sop", 'source: "https://example.com/v"', "tags:", '  - "sop"'):
            self.assertIn(key, head)
        self.assertIn("![畫面](attachments/step-0030.jpg)", md)
        self.assertNotIn("frames/", md)
        self.assertNotIn("frame@", md)
        self.assertTrue((d / "kb" / "attachments" / "step-0030.jpg").exists())
        self.assertIn("# 測試影片：操作 SOP", md)

    def test_sop_name_is_readable_and_filesystem_safe(self):
        self.assertEqual(br.sop_name({"title": "免費製作 AI 影片：ComfyUI 接上 Claude（MCP）"}),
                         "免費製作 AI 影片：ComfyUI 接上 Claude（MCP）_操作SOP")
        self.assertEqual(br.sop_name({"title": 'a\\b/c:d*e?f"g<h>i|j'}), "abcdefghij_操作SOP")
        self.assertEqual(br.sop_name({"title": "x" * 100}), "x" * 60 + "_操作SOP")

    def test_tags_have_no_spaces(self):
        d = self.make(self.SOP)
        import json
        m = json.loads((d / "meta.json").read_text(encoding="utf-8"))
        m["keywords"] = ["Comfy Cloud", "MCP"]
        (d / "meta.json").write_text(json.dumps(m, ensure_ascii=False), encoding="utf-8")
        br.build(d)
        head = (d / "kb" / "測試影片_操作SOP.md").read_text(encoding="utf-8").split("---\n")[1]
        self.assertIn('  - "Comfy-Cloud"', head)
        self.assertNotIn("Comfy Cloud", head)

    def test_no_sop_or_off_means_no_bundle_and_stale_bundle_removed(self):
        d = self.make(self.SOP)
        br.build(d)
        (d / "sop.md").unlink()
        br.build(d)
        self.assertFalse((d / "kb").exists())
        d2 = self.make(self.SOP, sop_mode="off")
        br.build(d2)
        self.assertFalse((d2 / "kb").exists())

    def test_rebuild_replaces_old_generated_attachments(self):
        d = self.make(self.SOP)
        br.build(d)
        (d / "kb" / "attachments" / "step-9999.jpg").write_bytes(b"x")
        br.build(d)
        self.assertFalse((d / "kb" / "attachments" / "step-9999.jpg").exists())

    def test_rebuild_keeps_users_own_files_in_kb(self):
        """kb/ may be opened as an Obsidian vault: .obsidian/ and the user's own notes/images must survive."""
        d = self.make(self.SOP)
        br.build(d)
        (d / "kb" / ".obsidian").mkdir()
        (d / "kb" / ".obsidian" / "app.json").write_text("{}", encoding="utf-8")
        (d / "kb" / "my note.md").write_text("mine", encoding="utf-8")
        (d / "kb" / "attachments" / "mine.png").write_bytes(b"x")
        br.build(d)
        self.assertTrue((d / "kb" / ".obsidian" / "app.json").exists())
        self.assertEqual((d / "kb" / "my note.md").read_text(encoding="utf-8"), "mine")
        self.assertTrue((d / "kb" / "attachments" / "mine.png").exists())
        self.assertTrue((d / "kb" / "測試影片_操作SOP.md").exists())
        (d / "sop.md").unlink()  # no SOP any more: only generated files go, user files stay
        br.build(d)
        self.assertFalse((d / "kb" / "測試影片_操作SOP.md").exists())
        self.assertTrue((d / "kb" / "my note.md").exists())
        self.assertTrue((d / "kb" / "attachments" / "mine.png").exists())

    def test_renamed_title_does_not_leave_old_sop_md_in_kb(self):
        import json
        d = self.make(self.SOP)
        br.build(d)
        m = json.loads((d / "meta.json").read_text(encoding="utf-8"))
        m["title"] = "新標題"
        (d / "meta.json").write_text(json.dumps(m, ensure_ascii=False), encoding="utf-8")
        br.build(d)
        self.assertEqual(self.kb_files(d), ["新標題_操作SOP.md"])

    def front(self, d):
        return (d / "kb" / "測試影片_操作SOP.md").read_text(encoding="utf-8").split("---\n")[1]

    def test_frontmatter_is_valid_yaml_for_hostile_values(self):
        import json
        try:
            import yaml
        except ImportError:
            self.skipTest("PyYAML not installed")
        d = self.make(self.SOP)
        m = json.loads((d / "meta.json").read_text(encoding="utf-8"))
        m.update({"title": "測試影片", "video_id": "123456", "url": "https://example.com/v?a=1&b=#x",
                  "keywords": ["*star", "#hash", "C++", "2026", "a: b", "[x]", "Comfy Cloud", "!bang", "中文"]})
        (d / "meta.json").write_text(json.dumps(m, ensure_ascii=False), encoding="utf-8")
        br.build(d)
        data = yaml.safe_load(self.front(d))
        self.assertEqual(data["video_id"], "123456")  # a string, not the integer 123456
        self.assertEqual(data["source"], "https://example.com/v?a=1&b=#x")
        self.assertTrue(all(isinstance(t, str) and t and not t.isdigit() for t in data["tags"]))
        self.assertIn("Comfy-Cloud", data["tags"])
        self.assertIn("中文", data["tags"])
        self.assertTrue(all(c not in t for t in data["tags"] for c in " #*!:[]+"))

    def test_frontmatter_title_with_backslash_and_quote(self):
        import json
        try:
            import yaml
        except ImportError:
            self.skipTest("PyYAML not installed")
        d = self.make(self.SOP)
        m = json.loads((d / "meta.json").read_text(encoding="utf-8"))
        m["title"] = 'C:\\new\\data "x"'
        (d / "meta.json").write_text(json.dumps(m, ensure_ascii=False), encoding="utf-8")
        br.build(d)
        f = [p for p in (d / "kb").glob("*.md")][0]
        data = yaml.safe_load(f.read_text(encoding="utf-8").split("---\n")[1])
        self.assertEqual(data["title"], 'C:\\new\\data "x"：操作 SOP')


class CleanFillers(unittest.TestCase):
    def setUp(self):
        import clean_fillers
        self.cf = clean_fillers

    def test_removes_ne_and_a_in_clauses(self):
        s = self.cf.strip_fillers
        self.assertEqual(s("所以這個叫至誠，如果至誠的話呢，你看"), "所以這個叫至誠，如果至誠的話，你看")
        self.assertEqual(s("那個在智者大師的書呢專以專實來解釋它"), "那個在智者大師的書專以專實來解釋它")
        self.assertEqual(s("起是什麼意思啊，起就是內心"), "起是什麼意思，起就是內心")
        self.assertEqual(s("什麼叫至啊？什麼叫誠啊？"), "什麼叫至？什麼叫誠？")

    def test_sentence_final_particles_only_at_clause_end(self):
        s = self.cf.strip_fillers
        self.assertEqual(s("好嘛，我們看下去"), "好，我們看下去")
        self.assertEqual(s("你看啦。"), "你看。")
        self.assertEqual(s("囉嗦的人"), "囉嗦的人")  # not clause-final: untouched
        self.assertEqual(s("啦啦隊"), "啦啦隊")

    def test_protected_words_survive(self):
        s = self.cf.strip_fillers
        for w in ("嘛呢叭咪吽", "呢喃自語", "阿彌陀佛", "啊彌陀佛", "啊難尊者", "啊賴耶識"):
            self.assertEqual(s(w), w)

    def test_no_leading_or_doubled_punctuation_left_behind(self):
        s = self.cf.strip_fillers
        self.assertEqual(s("啊，起就是內心"), "起就是內心")
        self.assertEqual(s("對啊，，你看"), "對，你看")
        self.assertEqual(s("是啊。"), "是。")
        self.assertEqual(s("好啊，。"), "好。")

    def test_idempotent_and_empty_result(self):
        s = self.cf.strip_fillers
        once = s("這個呢，你看啊，很好呢。")
        self.assertEqual(once, "這個，你看，很好。")
        self.assertEqual(s(once), once)
        self.assertEqual(s("啊"), "")
        self.assertEqual(s("呢，"), "")

    def test_clean_text_same_rules_as_file(self):
        raw = "_Source: whisper (groq)._\n\n```\n[00:01] 這個呢，很好啊。\n[00:05] 啊\n```\n"
        text, changed, dropped = self.cf.clean_text(raw)
        self.assertEqual(text, "_Source: whisper (groq)._\n\n```\n[00:01] 這個，很好。\n```\n")
        self.assertEqual((changed, dropped), (1, 1))
        self.assertEqual(self.cf.clean_text(text), (text, 0, 0))

    def test_file_only_touches_segment_lines_and_drops_empty_ones(self):
        d = Path(tempfile.mkdtemp())
        p = d / "transcript.md"
        p.write_text("# 呢啊 標題\n\n_Source: whisper (groq)._\n\n```\n[00:01] 這個呢，很好啊。\n[00:05] 啊\n[00:07] 嗯，好\n```\n\n## 備註\n- 呢啊 不動\n", encoding="utf-8")
        n = self.cf.clean_file(p)
        out = p.read_text(encoding="utf-8")
        self.assertEqual(out, "# 呢啊 標題\n\n_Source: whisper (groq)._\n\n```\n[00:01] 這個，很好。\n[00:07] 嗯，好\n```\n\n## 備註\n- 呢啊 不動\n")
        self.assertEqual(n, (1, 1))  # (lines changed, lines dropped)
        self.assertEqual(self.cf.clean_file(p), (0, 0))  # second run: nothing left to do


class FenceHelpers(unittest.TestCase):
    def test_language_tag_inside_fence_does_not_close_it(self):
        out = Path(tempfile.mkdtemp())
        (out / "supplement.md").write_text("## S1\n```text\n```python\nx = 1\n```\nafter\n", encoding="utf-8")
        lines = br.supplement_lines(out)
        i = lines.index("```")
        self.assertEqual(lines[i + 1:i + 3], ["```python", "x = 1"])
        self.assertEqual(lines[i + 3], "```")  # closed by the bare fence
        self.assertIn("after", lines)

    def test_unclosed_fence_in_summary_is_closed(self):
        md = "intro\n```\ncode with no end\n"
        fixed = br.close_open_fence(md)
        self.assertEqual(fixed.count("```"), 2)
        self.assertEqual(br.close_open_fence("fine\n```\nx\n```"), "fine\n```\nx\n```")

    def test_longer_fence_needs_equal_or_longer_close(self):
        self.assertFalse(br.fence_close("````", "```"))
        self.assertTrue(br.fence_close("````", "````"))
        self.assertTrue(br.fence_close("~~~", "~~~"))
        self.assertFalse(br.fence_close("```", "~~~"))


class DocxFenceAndLinks(unittest.TestCase):
    def build(self, md):
        from docx import Document  # noqa: F401
        d = Path(tempfile.mkdtemp())
        out = d / "t.docx"
        br.md_to_docx(d, md, out)
        from docx import Document as D
        return D(str(out))

    def test_fence_rendered_verbatim_monospace(self):
        doc = self.build("## x\n\n```\n# 模型路由\n- **Fable 5.1**：架構\n```\n")
        texts = [p.text for p in doc.paragraphs]
        self.assertIn("# 模型路由", texts)
        self.assertIn("- **Fable 5.1**：架構", texts)
        self.assertEqual([p.style.name for p in doc.paragraphs if p.text == "# 模型路由"], ["Normal"])  # not a heading

    def test_tilde_fence_in_docx(self):
        doc = self.build("## x\n\n~~~\n# not a heading\n~~~\n")
        self.assertEqual([p.style.name for p in doc.paragraphs if p.text == "# not a heading"], ["Normal"])

    def test_indented_bullet_is_a_bullet(self):
        doc = self.build("1. 步驟\n   - 子項目\n- 一般項目\n")
        styles = {p.text: p.style.name for p in doc.paragraphs}
        self.assertEqual(styles["子項目"], "List Bullet")
        self.assertEqual(styles["一般項目"], "List Bullet")
        self.assertEqual(styles["1. 步驟"], "Normal")

    def test_max_img_h_caps_landscape_image_height(self):
        from PIL import Image
        d = Path(tempfile.mkdtemp())
        (d / "frames").mkdir()
        Image.new("RGB", (320, 180), (200, 200, 200)).save(d / "frames" / "a.jpg")
        out = d / "t.docx"
        br.md_to_docx(d, "![x](frames/a.jpg)\n", out, max_img_h=6)
        from docx import Document as D
        self.assertLessEqual(D(str(out)).inline_shapes[0].height.cm, 6.01)

    def test_link_label_and_url(self):
        doc = self.build("- [說明](https://example.com/a_(b)) 與 [supplement.md](supplement.md)\n")
        xml = doc.element.xml
        self.assertIn("說明", xml)
        self.assertIn("supplement.md", xml)


class SopMode(unittest.TestCase):
    def test_default_is_auto(self):
        self.assertEqual(wa.resolve_sop_mode(False, False, ""), "auto")

    def test_flags_override_and_are_remembered(self):
        self.assertEqual(wa.resolve_sop_mode(True, False, ""), "on")
        self.assertEqual(wa.resolve_sop_mode(False, True, "on"), "off")
        self.assertEqual(wa.resolve_sop_mode(False, False, "off"), "off")  # rerun without flag keeps it

    def test_unknown_remembered_value_falls_back_to_auto(self):
        self.assertEqual(wa.resolve_sop_mode(False, False, "weird"), "auto")

    def test_both_flags_rejected_by_cli(self):
        import subprocess
        p = subprocess.run([sys.executable, str(Path(wa.__file__)), "https://youtu.be/x", "--sop", "--no-sop"],
                           capture_output=True, text=True, env=dict(os.environ, PYTHONUTF8="1"))
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("not allowed with", p.stderr)


class FillersMode(unittest.TestCase):
    def test_mode_resolution(self):
        self.assertEqual(wa.resolve_fillers_mode(False, False, ""), "auto")
        self.assertEqual(wa.resolve_fillers_mode(True, False, ""), "on")
        self.assertEqual(wa.resolve_fillers_mode(False, True, "on"), "off")
        self.assertEqual(wa.resolve_fillers_mode(False, False, "off"), "off")  # rerun keeps it
        self.assertEqual(wa.resolve_fillers_mode(False, False, "weird"), "auto")

    def test_auto_means_pure_land_lectures_only(self):
        self.assertTrue(wa.fillers_enabled("auto", ["common", "pure-land"]))
        self.assertTrue(wa.fillers_enabled("auto", ["pure-land"]))
        self.assertFalse(wa.fillers_enabled("auto", ["ai-video"]))
        self.assertFalse(wa.fillers_enabled("auto", []))
        self.assertTrue(wa.fillers_enabled("on", []))
        self.assertFalse(wa.fillers_enabled("off", ["pure-land"]))

    def test_auto_also_for_lecture_video_type(self):
        self.assertTrue(wa.fillers_enabled("auto", [], "lecture"))
        self.assertTrue(wa.fillers_enabled("auto", ["ai-video"], "lecture"))
        self.assertFalse(wa.fillers_enabled("auto", [], "tutorial"))
        self.assertFalse(wa.fillers_enabled("auto", [], ""))
        self.assertFalse(wa.fillers_enabled("off", ["pure-land"], "lecture"))  # --keep-fillers always wins
        self.assertTrue(wa.fillers_enabled("on", [], "tutorial"))
        t = "```\n[00:01] 這個呢，很好啊。\n```\n"
        self.assertEqual(wa.apply_fillers(t, True, "auto", [], "lecture")[1:], (1, 0))

    def test_apply_fillers_only_for_chinese_and_enabled_mode(self):
        t = "```\n[00:01] 這個呢，很好啊。\n```\n"
        cleaned = "```\n[00:01] 這個，很好。\n```\n"
        self.assertEqual(wa.apply_fillers(t, True, "auto", ["pure-land"]), (cleaned, 1, 0))
        self.assertEqual(wa.apply_fillers(t, True, "on", []), (cleaned, 1, 0))
        self.assertEqual(wa.apply_fillers(t, False, "on", ["pure-land"]), (t, 0, 0))  # ja / ko / unclear text untouched
        self.assertEqual(wa.apply_fillers(t, True, "auto", ["ai-video"]), (t, 0, 0))
        self.assertEqual(wa.apply_fillers(t, True, "off", ["pure-land"]), (t, 0, 0))

    def test_both_flags_rejected_by_cli(self):
        import subprocess
        p = subprocess.run([sys.executable, str(Path(wa.__file__)), "https://youtu.be/x", "--clean-fillers", "--keep-fillers"],
                           capture_output=True, text=True, env=dict(os.environ, PYTHONUTF8="1"))
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("not allowed with", p.stderr)


class PickWhisper(unittest.TestCase):
    def test_groq_when_key_exists(self):
        self.assertEqual(wa.pick_whisper(None, False, "", True), "groq")

    def test_no_key_defers_to_watch_config(self):
        self.assertIsNone(wa.pick_whisper(None, False, "", False))

    def test_local_switch_beats_groq_key(self):
        self.assertEqual(wa.pick_whisper(None, True, "", True), "whisperx")

    def test_explicit_backend_wins(self):
        self.assertEqual(wa.pick_whisper("openai", True, "whisperx", True), "openai")

    def test_remembered_choice_beats_groq_key(self):
        self.assertEqual(wa.pick_whisper(None, False, "whisperx", True), "whisperx")

    def test_explicit_overrides_remembered(self):
        self.assertEqual(wa.pick_whisper("groq", False, "whisperx", True), "groq")


class HasGroqKey(unittest.TestCase):
    def check(self, env_value, file_text):
        with tempfile.TemporaryDirectory() as d:
            cfg = Path(d) / ".env"
            cfg.write_text(file_text, encoding="utf-8")
            env = {k: v for k, v in os.environ.items() if k != "GROQ_API_KEY"}
            if env_value is not None:
                env["GROQ_API_KEY"] = env_value
            with mock.patch.dict(os.environ, env, clear=True), mock.patch.object(wa.watch_config, "CONFIG_FILE", cfg), \
                    mock.patch.object(wa.watch_config.Path, "cwd", return_value=Path(d) / "nowhere"):
                return wa.has_groq_key()

    def test_key_in_env(self):
        self.assertTrue(self.check("k", ""))

    def test_key_only_in_config_file(self):
        self.assertTrue(self.check(None, 'GROQ_API_KEY="k"\n'))

    def test_blank_everywhere(self):
        self.assertFalse(self.check("  ", "GROQ_API_KEY=\n"))


class WhisperHints(unittest.TestCase):
    def test_api_backend_gets_language_and_prompt(self):
        env, lang, prompt, warnings = wa.whisper_hints("groq", "zh", "詞1、詞2")
        self.assertEqual(env, {"WATCH_WHISPER_LANGUAGE": "zh", "WATCH_WHISPER_PROMPT": "詞1、詞2"})
        self.assertEqual((lang, prompt, warnings), ("zh", "詞1、詞2", []))

    def test_api_backend_without_hints_sets_nothing(self):
        self.assertEqual(wa.whisper_hints("groq", "", ""), ({}, "", "", []))

    def test_whisperx_maps_language_and_drops_prompt(self):
        env, lang, prompt, warnings = wa.whisper_hints("whisperx", "zh", "詞1")
        self.assertEqual(env, {"WATCH_WHISPERX_LANGUAGE": "zh"})
        self.assertEqual((lang, prompt), ("zh", ""))
        self.assertEqual(len(warnings), 1)

    def test_whisperx_rejects_non_two_letter_language(self):
        env, lang, _, warnings = wa.whisper_hints("whisperx", "zh-TW", "")
        self.assertEqual((env, lang), ({}, ""))
        self.assertEqual(len(warnings), 1)


class Category(unittest.TestCase):
    CATS = ["AI", "佛學", "未分類"]

    def pick(self, arg="", prev="", domains=()):
        return wa.pick_category(arg, prev, list(domains), self.CATS)

    def test_precedence_arg_prev_domain_default(self):
        self.assertEqual(self.pick("佛學", "AI", ["ai-chips"]), "佛學")
        self.assertEqual(self.pick("", "AI", ["pure-land"]), "AI")
        self.assertEqual(self.pick("", "", ["x", "pure-land"]), "佛學")
        self.assertEqual(self.pick(), wa.DEFAULT_CATEGORY)

    def test_unknown_category_is_error_but_ascii_case_is_fixed(self):
        with self.assertRaises(ValueError):
            self.pick("理財")
        self.assertEqual(self.pick("ai"), "AI")
        with self.assertRaises(ValueError):
            self.pick("佛教")

    @unittest.skipIf(wa._S2T is None, "opencc not installed")
    def test_simplified_category_becomes_traditional(self):
        self.assertEqual(self.pick("佛学"), "佛學")
        self.assertEqual(self.pick("未分类"), "未分類")

    def test_stale_remembered_category_asks_for_category(self):
        with self.assertRaisesRegex(ValueError, "--category"):
            self.pick("", "舊領域")

    def test_first_glossary_domain_wins_on_disagreement(self):
        self.assertEqual(self.pick("", "", ["pure-land", "ai-chips"]), "佛學")
        self.assertEqual(self.pick("", "", ["ai-chips", "pure-land"]), "AI")

    def test_every_glossary_domain_category_is_allowed(self):
        cats = wa.load_categories()
        self.assertEqual(cats[-1], wa.DEFAULT_CATEGORY)
        for f in (wa.ROOT / "glossary" / "domains").glob("*.json"):
            c = __import__("json").loads(f.read_text(encoding="utf-8")).get("category")
            self.assertIn(c, cats, f.name)

    @unittest.skipIf(wa._S2T is None, "opencc not installed")
    def test_category_names_are_traditional_chinese(self):
        for c in wa.load_categories():
            self.assertEqual(wa.to_traditional(c), c, f"category '{c}' must be Traditional Chinese")

    def test_index_sections_in_category_order_newest_first(self):
        import json
        with tempfile.TemporaryDirectory() as t, mock.patch.object(wa, "OUT_ROOT", Path(t)):
            for cat, day, vid in [("佛學", "2026-01-01", "a"), ("AI", "2026-01-01", "b"), ("AI", "2026-02-01", "c")]:
                d = Path(t) / cat / f"{day}_X_t__{vid}"
                d.mkdir(parents=True)
                (d / "meta.json").write_text(json.dumps({"last_updated": day, "platform": "X", "title": vid, "duration": "0", "transcript_source": "s"}), encoding="utf-8")
            wa.update_index()
            idx = (Path(t) / "index.md").read_text(encoding="utf-8")
        self.assertLess(idx.index("## AI（2）"), idx.index("## 佛學（1）"))
        self.assertLess(idx.index("| c |"), idx.index("| b |"))
        self.assertNotIn("未分類", idx)

    def test_safe_category_removed(self):
        self.assertFalse(hasattr(wa, "safe_category"))

    def test_nested_category_allowed_and_indexed_after_parent(self):
        import json
        cats = ["AI", "AI/claude-code", "佛學", "未分類"]
        self.assertEqual(wa.pick_category("AI/claude-code", "", [], cats), "AI/claude-code")
        self.assertEqual(wa.pick_category("ai/CLAUDE-CODE", "", [], cats), "AI/claude-code")
        self.assertEqual(wa.pick_category("", "", ["claude-code"], wa.load_categories()), "AI/claude-code")
        with tempfile.TemporaryDirectory() as t, mock.patch.object(wa, "OUT_ROOT", Path(t)):
            for cat, vid in [("AI", "a"), ("AI/claude-code", "b")]:
                d = Path(t) / cat / f"2026-01-01_X_t__{vid}"
                d.mkdir(parents=True)
                (d / "meta.json").write_text(json.dumps({"last_updated": "2026-01-01", "platform": "X", "title": vid, "duration": "0", "transcript_source": "s"}), encoding="utf-8")
            self.assertEqual([p.parent.name for p in wa.find_dirs("b")], ["claude-code"])
            wa.update_index()
            idx = (Path(t) / "index.md").read_text(encoding="utf-8")
        self.assertLess(idx.index("## AI（1）"), idx.index("## AI／claude-code（1）"))
        self.assertIn("(./AI/claude-code/2026-01-01_X_t__b/)", idx)

    def test_find_dirs_nested_and_legacy(self):
        with tempfile.TemporaryDirectory() as t, mock.patch.object(wa, "OUT_ROOT", Path(t)):
            (Path(t) / "AI" / "2026-01-01_X_a__v1").mkdir(parents=True)
            (Path(t) / "2026-01-02_X_b__v2").mkdir()
            self.assertEqual([p.name for p in wa.find_dirs("v1")], ["2026-01-01_X_a__v1"])
            self.assertEqual([p.name for p in wa.find_dirs("v2")], ["2026-01-02_X_b__v2"])


class MindMap(unittest.TestCase):
    SLIDES = [{"start": "00:00", "frame": "cue_0001.jpg", "title": "開場", "keywords": ["甲", "乙"]},
              {"start": "00:30", "frame": "cue_0001.jpg", "title": "第一步", "keywords": []}]

    def make(self, slides=None, mindmap_md=None):
        import json
        from PIL import Image
        d = Path(tempfile.mkdtemp()) / "2026-10-21_Web_測試__abc"
        (d / "frames").mkdir(parents=True)
        Image.new("RGB", (320, 180), (200, 200, 200)).save(d / "frames" / "cue_0001.jpg")
        meta = {"url": "https://example.com/v", "platform": "Web", "title": "測試影片", "duration": "01:00 (60s)",
                "slides": slides if slides is not None else self.SLIDES}
        (d / "meta.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
        (d / "summary.md").write_text("# 摘要\n\n內容\n", encoding="utf-8")
        (d / "transcript.md").write_text("```\n[00:01] 你好\n[00:31] 第一步\n```\n", encoding="utf-8")
        (d / "frames.json").write_text(json.dumps([{"file": "cue_0001.jpg", "t": "00:00"}]), encoding="utf-8")
        if mindmap_md is not None:
            (d / "mindmap.md").write_text(mindmap_md, encoding="utf-8")
        return d

    def test_parse_outline_single_root_and_several_tops(self):
        mm = sys.modules.get("mindmap") or __import__("mindmap")
        self.assertEqual(mm.parse_outline("# t\n- 中心\n  - 分支\n    - 葉\n"), ("中心", [("分支", [("葉", [])])]))
        self.assertEqual(mm.parse_outline("- a\n- b\n", "標題"), ("標題", [("a", []), ("b", [])]))
        self.assertIsNone(mm.parse_outline("沒有條列"))

    def test_auto_from_slides_then_regenerated_but_hand_written_kept(self):
        d = self.make()
        br.build(d)
        auto = (d / "mindmap.md").read_text(encoding="utf-8")
        self.assertTrue(auto.startswith(__import__("mindmap").AUTO_MARK))
        self.assertIn("  - 開場", auto)
        self.assertIn("    - 甲", auto)
        self.assertTrue((d / "mindmap.png").exists())
        (d / "mindmap.md").write_text("- 自己寫的中心\n  - 分支\n", encoding="utf-8")
        br.build(d)
        self.assertEqual((d / "mindmap.md").read_text(encoding="utf-8"), "- 自己寫的中心\n  - 分支\n")

    def test_section_before_frames_section_with_links(self):
        d = self.make()
        md = br.build(d)[0].read_text(encoding="utf-8")
        self.assertLess(md.index("## 摘要與名詞說明"), md.index("## 思維導圖"))
        self.assertLess(md.index("## 思維導圖"), md.index("## 圖文對照"))
        self.assertIn("![思維導圖](mindmap.png)", md)
        self.assertIn("[mindmap.md](mindmap.md)", md.split("## 補充資料連結")[1])

    def test_docx_links_to_local_files_and_shows_map(self):
        from docx import Document
        d = self.make()
        outs = br.build(d)
        doc = Document(str(outs[1]))
        targets = {r.target_ref for r in doc.part.rels.values() if r.reltype.endswith("/hyperlink")}
        self.assertIn("mindmap.png", targets)
        self.assertIn("mindmap.md", targets)
        self.assertEqual(len(doc.inline_shapes), 3)  # mind map + one frame per slide

    def test_no_slides_and_no_outline_means_no_map(self):
        d = self.make(slides=[])
        md = br.build(d)[0].read_text(encoding="utf-8")
        self.assertNotIn("思維導圖", md)
        self.assertFalse((d / "mindmap.png").exists())
        self.assertFalse((d / "mindmap.md").exists())


class ReviewLinks(unittest.TestCase):
    def test_repo_slug_from_common_remotes(self):
        for url in ("https://github.com/o/r.git", "git@github.com:o/r.git", "http://local_proxy@127.0.0.1:1/git/o/r"):
            self.assertEqual(rl.repo_slug(url), "o/r")

    def test_links_are_percent_encoded_and_point_at_branch(self):
        links = dict(rl.build_links("o/r", "claude/x", "watch-output/AI/a b", ["summary.md"]))
        self.assertEqual(links["資料夾"], "https://github.com/o/r/tree/claude/x/watch-output/AI/a%20b")
        self.assertEqual(links["summary.md"], "https://github.com/o/r/blob/claude/x/watch-output/AI/a%20b/summary.md")
        self.assertTrue(links["與 main 的差異（開 PR 前看這個）"].endswith("/compare/main...claude/x"))


class PickFrames(unittest.TestCase):
    def test_windows_run_from_each_start_to_the_next(self):
        self.assertEqual(pf.windows(["00:00", "00:11", "01:05"], 80), [(0, 11), (11, 65), (65, 81)])

    def test_best_skips_boundary_seconds_when_window_is_long(self):
        score = [9, 9, 1, 5, 1, 1, 1, 9]  # seconds 0-1 still show the previous slide, 7 is the next transition
        self.assertEqual(pf.best_in_window(score, 0, 8), 3)

    def test_best_uses_whole_short_window_and_prefers_later_tie(self):
        self.assertEqual(pf.best_in_window([2, 5, 5], 0, 3), 2)

    def test_spoken_lines_only_inside_window(self):
        tr = "# t\n```\n[00:09] a\n[00:10] b\n[00:19] c\n[00:20] d\n```"
        self.assertEqual(pf.spoken_lines(tr, 10, 20), ["[00:10] b", "[00:19] c"])

    def test_register_slides_points_meta_and_frames_at_slide_files(self):
        import json
        with tempfile.TemporaryDirectory() as t:
            d = Path(t)
            (d / "meta.json").write_text(json.dumps({"slides": [{"start": "00:00"}, {"start": "00:20"}]}), encoding="utf-8")
            (d / "frames.json").write_text(json.dumps([{"file": "frame_0001.jpg", "t": "00:00", "reason": "x"},
                                                       {"file": "slide_09.jpg", "t": "00:09", "reason": "clean-slide"}]), encoding="utf-8")
            pf.register_slides(d, [3, 31])
            meta = json.loads((d / "meta.json").read_text(encoding="utf-8"))
            frames = json.loads((d / "frames.json").read_text(encoding="utf-8"))
        self.assertEqual([s["frame"] for s in meta["slides"]], ["slide_01.jpg", "slide_02.jpg"])
        self.assertEqual([(f["file"], f["t"]) for f in frames if f["reason"] == "clean-slide"],
                         [("slide_01.jpg", "00:03"), ("slide_02.jpg", "00:31")])
        self.assertIn("frame_0001.jpg", [f["file"] for f in frames])


class WatchArchiveSkillDocs(unittest.TestCase):
    """The skill outline stays small enough to survive auto-compact, and every pointer to its reference files resolves."""
    SKILL_DIR = wa.ROOT / ".claude" / "skills" / "watch-archive"
    REFS = SKILL_DIR / "references"
    SKILL_MAX_CHARS = 6000

    def skill_text(self):
        return (self.SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")

    def test_outline_table_lists_exactly_the_reference_files(self):
        listed = {m for line in self.skill_text().splitlines() if line.startswith("|")
                  for m in re.findall(r"`([\w.-]+\.md)`", line)}
        self.assertEqual(listed, {p.name for p in self.REFS.glob("*.md")})

    def test_outline_is_under_size_limit(self):
        self.assertLessEqual(len(self.skill_text()), self.SKILL_MAX_CHARS)

    def test_reference_paths_outside_the_skill_resolve_from_repo_root(self):
        full = ".claude/skills/watch-archive/references/"
        for f in [wa.ROOT / "CLAUDE.md", wa.ROOT / "README.md", *sorted((wa.ROOT / "scripts").glob("*.py"))]:
            for m in re.finditer(r"[A-Za-z0-9_./-]*references/[A-Za-z0-9_.-]+\.md", f.read_text(encoding="utf-8")):
                path = m.group(0)
                with self.subTest(file=f.name, path=path):
                    self.assertTrue(path.startswith(full), "write the full path from the repo root")
                    self.assertTrue((wa.ROOT / path).is_file())

    def test_skill_files_name_other_reference_files_by_full_path(self):
        """Only the outline's table may use bare file names; everywhere else a reader starting at the repo root must find the file."""
        names = "|".join(re.escape(p.name) for p in self.REFS.glob("*.md"))
        for f in [self.SKILL_DIR / "SKILL.md", *sorted(self.REFS.glob("*.md"))]:
            for line in f.read_text(encoding="utf-8").splitlines():
                if f.name == "SKILL.md" and line.startswith("|"):
                    continue
                for m in re.finditer(rf"(?<![A-Za-z0-9_./-])({names})", line):
                    with self.subTest(file=f.name, name=m.group(1)):
                        self.fail(f"write .claude/skills/watch-archive/references/{m.group(1)}")

    def test_archive_command_templates_do_not_force_a_language(self):
        for f in sorted(self.REFS.glob("*.md")):
            for line in f.read_text(encoding="utf-8").splitlines():
                if "scripts/watch_archive.py" in line:
                    with self.subTest(file=f.name):
                        self.assertNotIn("--lang", line, "the archive template must not force --lang (CLAUDE.md hard rule)")


if __name__ == "__main__":
    unittest.main()
